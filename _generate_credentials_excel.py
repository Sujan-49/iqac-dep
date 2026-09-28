import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
# remove default sheet
wb.remove(wb.active)

# Helper function to style a worksheet
def create_credential_sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title=title)
    ws.views.sheetView[0].showGridLines = True

    # Colors
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Royal Blue
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Calibri", size=11)
    bold_font = Font(name="Calibri", size=11, bold=True)
    
    thin_border = Border(
        left=Side(style='thin', color='D4D4D4'),
        right=Side(style='thin', color='D4D4D4'),
        top=Side(style='thin', color='D4D4D4'),
        bottom=Side(style='thin', color='D4D4D4')
    )

    # Write headers
    ws.append(headers)
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    
    ws.row_dimensions[1].height = 24

    # Write rows
    for row_num, row_data in enumerate(rows, 2):
        ws.append(row_data)
        ws.row_dimensions[row_num].height = 20
        for col_num, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.font = cell_font
            cell.border = thin_border
            if col_num in [2, 3, 4, 7, 8]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            if col_num in [2, 3]:
                cell.font = bold_font

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

headers = ["Name", "Username", "Password", "Role", "Department", "Allocated Subject", "Batch", "Status"]

# 1. Students DI (60 students)
firstNamesMale = ['Aravind', 'Bala', 'Chandran', 'Dinesh', 'Gokul', 'Karthik', 'Lokesh', 'Manoj', 'Naveen', 'Pradeep', 'Rahul', 'Santhosh', 'Vignesh', 'Vijay', 'Yogesh', 'Surya', 'Pravin', 'Siddharth', 'Vishal', 'Ashwin']
firstNamesFemale = ['Ananya', 'Bhavana', 'Divya', 'Harini', 'Kavya', 'Keerthana', 'Meena', 'Nivetha', 'Priya', 'Ramya', 'Sruthi', 'Swetha', 'Vaishnavi', 'Yamini', 'Abirami', 'Janani', 'Madhumitha', 'Pavithra', 'Subhashini', 'Vidyashree']
lastNames = ['Subramanian', 'Ramesh', 'Kumar', 'Senthil', 'Murugan', 'Natarajan', 'Balakrishnan', 'Venkatesan', 'Sundaram', 'Ganesan', 'Elango', 'Manoharan', 'Chellappa', 'Viswanathan']

def get_name(i):
    is_female = (i % 2 == 0)
    first = firstNamesFemale[i % len(firstNamesFemale)] if is_female else firstNamesMale[i % len(firstNamesMale)]
    last = lastNames[(i * 3) % len(lastNames)]
    return f"{first} {last}"

rows_std_di = []
for i in range(1, 61):
    u = f"24DI{i:02d}"
    rows_std_di.append([get_name(i), u, u, "Student", "Diploma Information Technology", "-", "24DI", "Approved"])

create_credential_sheet(wb, "Students_DI", headers, rows_std_di)

# 2. Students AI (60 students)
rows_std_ai = []
for i in range(1, 61):
    u = f"24AI{i:02d}"
    rows_std_ai.append([get_name(i + 60), u, u, "Student", "Artificial Intelligence & Machine Learning", "-", "24AI", "Approved"])

create_credential_sheet(wb, "Students_AI", headers, rows_std_ai)

# 3. Teachers DI
rows_teachers_di = [
    ["Prof. R. Ramesh", "tutor_di", "staff123", "Tutor", "Diploma Information Technology", "Batch Monitoring (24DI)", "24DI", "Active"],
    ["Prof. A. Priya", "wt_di", "staff123", "Teacher", "Diploma Information Technology", "DI501 - Web Technology", "24DI", "Active"],
    ["Prof. B. Vijay", "crypto_di", "staff123", "Teacher", "Diploma Information Technology", "DI502 - Cryptography & Security", "24DI", "Active"],
    ["Prof. C. Anand", "cloud_di", "staff123", "Teacher", "Diploma Information Technology", "DI503 - Cloud Computing", "24DI", "Active"],
    ["Prof. D. Divya", "ai_di", "staff123", "Teacher", "Diploma Information Technology", "DI504 - Artificial Intelligence", "24DI", "Active"],
    ["Prof. E. Karthik", "lab_di", "staff123", "Teacher", "Diploma Information Technology", "DI505 - Web Technology Lab", "24DI", "Active"],
]
create_credential_sheet(wb, "Teachers_DI", headers, rows_teachers_di)

# 4. Teachers AI
rows_teachers_ai = [
    ["Prof. S. Suresh", "tutor_ai", "staff123", "Tutor", "Artificial Intelligence & Machine Learning", "Batch Monitoring (24AI)", "24AI", "Active"],
    ["Prof. F. Kavitha", "python_ai", "staff123", "Teacher", "Artificial Intelligence & Machine Learning", "AI501 - Python for AI", "24AI", "Active"],
    ["Prof. G. Rajesh", "ml_ai", "staff123", "Teacher", "Artificial Intelligence & Machine Learning", "AI502 - Machine Learning", "24AI", "Active"],
    ["Prof. H. Deepika", "dl_ai", "staff123", "Teacher", "Artificial Intelligence & Machine Learning", "AI503 - Deep Learning", "24AI", "Active"],
    ["Prof. I. Manoj", "cv_ai", "staff123", "Teacher", "Artificial Intelligence & Machine Learning", "AI504 - Computer Vision", "24AI", "Active"],
    ["Prof. J. Nithya", "ailab_ai", "staff123", "Teacher", "Artificial Intelligence & Machine Learning", "AI505 - AI Laboratory", "24AI", "Active"],
]
create_credential_sheet(wb, "Teachers_AI", headers, rows_teachers_ai)

# 5. HOD
rows_hod = [
    ["Dr. K. Senthil Kumar", "hod_di", "hod123", "HOD", "Diploma Information Technology", "Department Oversight", "All", "Active"],
    ["Dr. M. Arunkumar", "hod_ai", "hod123", "HOD", "Artificial Intelligence & Machine Learning", "Department Oversight", "All", "Active"],
]
create_credential_sheet(wb, "HOD", headers, rows_hod)

# 6. IQAC
rows_iqac = [
    ["IQAC Coordinator", "iqac", "iqac123", "IQAC", "Institutional Quality Assurance Cell", "All Department Reports & NBA Audit", "All", "Active"]
]
create_credential_sheet(wb, "IQAC", headers, rows_iqac)

# 7. Principal
rows_principal = [
    ["Principal", "principal", "principal123", "Principal", "PSG Polytechnic College", "Institutional Administration", "All", "Active"]
]
create_credential_sheet(wb, "Principal", headers, rows_principal)

# 8. Super Admin
rows_admin = [
    ["System Administrator", "admin", "admin123", "Super Admin", "Central Administration", "Full System Access", "All", "Active"]
]
create_credential_sheet(wb, "Super_Admin", headers, rows_admin)

excel_filename = "c:/xampp/htdocs/Iqac/ERP_Login_Credentials.xlsx"
wb.save(excel_filename)
print(f"Successfully generated Excel workbook: {excel_filename}")
