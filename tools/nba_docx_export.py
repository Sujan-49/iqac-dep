import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def clear_body(doc):
    body = doc._body._element
    for child in list(body):
        if child.tag.endswith('sectPr'):
            continue
        body.remove(child)


def set_cell_text(cell, value, bold=False):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(value))
    run.bold = bold
    run.font.size = Pt(8)


def add_heading(doc, text, size=12):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    return p


def add_meta_table(doc, meta):
    pairs = list(meta.items())
    table = doc.add_table(rows=0, cols=4)
    table.style = 'Table Grid'
    for i in range(0, len(pairs), 2):
        row = table.add_row().cells
        k1, v1 = pairs[i]
        set_cell_text(row[0], k1, True)
        set_cell_text(row[1], v1)
        if i + 1 < len(pairs):
            k2, v2 = pairs[i + 1]
            set_cell_text(row[2], k2, True)
            set_cell_text(row[3], v2)
        else:
            set_cell_text(row[2], '')
            set_cell_text(row[3], '')
    doc.add_paragraph()


def add_data_table(doc, title, headers, rows):
    p = doc.add_paragraph()
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(10)
    table = doc.add_table(rows=1, cols=max(1, len(headers)))
    table.style = 'Table Grid'
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, True)
    for row_data in rows:
        cells = table.add_row().cells
        for i, cell in enumerate(cells):
            set_cell_text(cell, row_data[i] if i < len(row_data) else '')
    doc.add_paragraph()


def add_signature_area(doc):
    add_data_table(
        doc,
        'Signature Area',
        ['Prepared By', 'Faculty In-charge', 'Tutor', 'HOD', 'IQAC Coordinator', 'Principal'],
        [['', '', '', '', '', '']],
    )
    seal = doc.add_paragraph()
    seal.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = seal.add_run('Department Seal Area')
    run.bold = True
    run.font.size = Pt(10)


def main():
    if len(sys.argv) != 4:
        raise SystemExit('Usage: nba_docx_export.py input.json template.docx output.docx')
    data = json.loads(Path(sys.argv[1]).read_text(encoding='utf-8-sig'))
    template = Path(sys.argv[2])
    output = Path(sys.argv[3])
    doc = Document(str(template))
    clear_body(doc)

    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)

    add_heading(doc, 'PSG Polytechnic College', 14)
    add_heading(doc, 'Internal Quality Assurance Cell (IQAC) Formats and Procedures', 11)
    add_heading(doc, 'NBA ACCREDITATION REPORT', 11)
    add_heading(doc, data.get('title', 'NBA REPORT'), 10)
    doc.add_paragraph()
    add_meta_table(doc, data.get('meta', {}))

    for section_data in data.get('sections', []):
        add_data_table(
            doc,
            section_data.get('title', 'Report Table'),
            section_data.get('headers', []),
            section_data.get('rows', []),
        )
    add_signature_area(doc)
    doc.save(str(output))


if __name__ == '__main__':
    main()
